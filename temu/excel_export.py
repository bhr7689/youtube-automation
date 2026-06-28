"""테무 대량 등록용 엑셀 생성기.

가이드 Part2 의 대량 등록 흐름:
  셀러센터 → 제품추가 → 대량 템플릿 → 카테고리 선택 → 스프레드시트 다운로드
  → (여기서 채워) → 스프레드시트 업로드

실제 셀러센터가 발급하는 컬럼은 카테고리별로 다르므로,
이 파일은 '범용 보조 엑셀'을 생성한다:
  Sheet 'Listing'    : 단일 행 — 제목/설명/속성/배송/ITC
  Sheet 'SKU'        : 변형별 행 (색상/사이즈/SKU코드/가격/수량/치수/이미지파일명)
  Sheet 'Attributes' : 키-값 속성 (재질/스타일/관리법…)
사용자는 셀러센터에서 받은 공식 템플릿에 이 엑셀의 컬럼을 그대로 복붙하면 된다.
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .schema import TemuListing


HEADER_FILL = PatternFill("solid", fgColor="FFEDD5")
HEADER_FONT = Font(bold=True, color="9A3412")
WRAP = Alignment(wrap_text=True, vertical="top")


def _write_header(ws, row: int, headers: list[str]) -> None:
    for col_i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col_i, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autosize(ws, max_width: int = 60) -> None:
    for col in ws.columns:
        first = next(iter(col))
        col_letter = get_column_letter(first.column)
        longest = 0
        for cell in col:
            v = cell.value
            if v is None:
                continue
            for line in str(v).splitlines():
                longest = max(longest, len(line))
        ws.column_dimensions[col_letter].width = min(max(longest + 2, 10), max_width)


def build_excel(listing: TemuListing, *, image_manifest: dict | None = None) -> bytes:
    wb = Workbook()

    # 1) Listing 시트
    ws1 = wb.active
    ws1.title = "Listing"
    headers = [
        "카테고리", "브랜드", "제목(500자)", "상세 설명", "ITC 코드",
        "처리 시간(일)", "배송 템플릿", "배송 방법",
        "메인 이미지", "갤러리 이미지(콤마)", "동영상",
        "옵션축1", "옵션축2",
        "안전·규정 메모", "규정 첨부(콤마)",
    ]
    _write_header(ws1, 1, headers)
    main_img = (image_manifest or {}).get("main", "") if image_manifest else listing.main_image_filename
    gallery = (image_manifest or {}).get("gallery", []) if image_manifest else listing.gallery_filenames
    ws1.append([
        listing.category,
        listing.brand_name,
        listing.final_title(),
        listing.description,
        listing.itc_code,
        listing.shipping.processing_days,
        listing.shipping.template_name,
        listing.shipping.method,
        main_img,
        ", ".join(gallery),
        listing.video_filename,
        listing.option_axis_1,
        listing.option_axis_2,
        listing.compliance_notes,
        ", ".join(listing.compliance_filenames),
    ])
    for cell in ws1[2]:
        cell.alignment = WRAP
    ws1.row_dimensions[2].height = 120
    _autosize(ws1)

    # 2) SKU 시트
    ws2 = wb.create_sheet("SKU")
    sku_headers = [
        listing.option_axis_1 or "옵션1",
        listing.option_axis_2 or "옵션2",
        "SKU 코드", "수량", "기본가격(KRW)",
        "무게(g)", "가로(cm)", "세로(cm)", "높이(cm)",
        "SKU 이미지",
    ]
    _write_header(ws2, 1, sku_headers)
    sku_map = (image_manifest or {}).get("sku", {}) if image_manifest else {}
    for sku in listing.skus:
        key = f"{sku.color}_{sku.size}".strip("_")
        ws2.append([
            sku.color, sku.size, sku.sku_code, sku.qty, sku.price_krw,
            sku.weight_g, sku.length_cm, sku.width_cm, sku.height_cm,
            sku_map.get(key, sku.image_filename),
        ])
    _autosize(ws2)

    # 3) Attributes 시트
    ws3 = wb.create_sheet("Attributes")
    _write_header(ws3, 1, ["속성명", "값"])
    for a in listing.attributes:
        ws3.append([a.label, a.value])
    _autosize(ws3, max_width=80)

    # 4) ReadMe 시트
    ws4 = wb.create_sheet("README")
    notes = [
        ["테무 한국 셀러센터 — 대량 등록용 보조 엑셀"],
        [""],
        ["■ 사용 순서"],
        ["  1) 셀러센터 → 제품 → 제품추가 → 대량 템플릿"],
        ["  2) 카테고리 선택 → 스프레드시트 생성 → 스프레드시트 다운로드"],
        ["  3) 이 엑셀의 'Listing'·'SKU'·'Attributes' 값을 공식 템플릿에 복사"],
        ["  4) 이미지 ZIP 안의 파일들을 셀러센터 폼에 업로드"],
        [""],
        ["■ 제목 권장 형식 (가이드 1단계)"],
        ["  [브랜드] + [세부정보] + [적용범위] + [제품유형] + [주요특징/기능/장점]"],
        ["  최대 500자"],
        [""],
        ["■ ITC 코드"],
        ["  Gen Standard  — 과세 상품 (대부분)"],
        ["  Gen Exempt    — 영세율·면세 상품"],
    ]
    for row in notes:
        ws4.append(row)
    ws4.column_dimensions["A"].width = 70

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
